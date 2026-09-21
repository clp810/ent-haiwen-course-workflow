#!/usr/bin/env python3
"""Build a consistently styled coursework DOCX from a compact JSON specification."""

from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)
RED = RGBColor(156, 0, 6)
BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
GRID = "D9D9D9"
ALLOWED_BLOCKS = {"paragraph", "bullets", "table", "sources", "page_break"}
FINAL_BLOCKERS = re.compile(
    r"\b(?:draft|pending|todo|tbc|tbd)\b|待确认|待补充|\{\{|\}\}|<placeholder>",
    re.IGNORECASE,
)


def load_spec(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        spec = json.load(handle)
    if not isinstance(spec, dict):
        raise ValueError("Document spec must be a JSON object.")
    for key in ("title", "stage", "sections"):
        if key not in spec:
            raise ValueError(f"Document spec is missing {key!r}.")
    if spec["stage"] not in {"DRAFT", "FINAL"}:
        raise ValueError("stage must be DRAFT or FINAL.")
    if not isinstance(spec["sections"], list):
        raise ValueError("sections must be an array.")
    for section in spec["sections"]:
        if not isinstance(section, dict) or not isinstance(section.get("blocks", []), list):
            raise ValueError("Each section must be an object with a blocks array.")
        for block in section.get("blocks", []):
            if block.get("type") not in ALLOWED_BLOCKS:
                raise ValueError(f"Unsupported block type: {block.get('type')!r}")
    serialized = json.dumps(spec, ensure_ascii=False)
    if spec["stage"] == "FINAL" and FINAL_BLOCKERS.search(serialized):
        raise ValueError("FINAL spec contains a draft, pending, or placeholder marker.")
    return spec


def set_font(run, size=11, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = properties.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        properties.append(shade)
    shade.set(qn("w:fill"), fill)


def set_margins(cell, value=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)


def set_width(cell, inches: float) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def configure(doc: Document, spec: dict) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        border = style._element.get_or_add_pPr().find(qn("w:pBdr"))
        if border is not None:
            style._element.get_or_add_pPr().remove(border)

    properties = spec.get("properties", {})
    author = properties.get("author", "Course Team")
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    doc.core_properties.subject = properties.get("subject", "Coursework")
    doc.core_properties.category = spec["stage"]


def add_metadata(doc: Document, metadata: list[dict]) -> None:
    if not metadata:
        return
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    for index, item in enumerate(metadata):
        label_cell, value_cell = table.rows[index].cells
        set_width(label_cell, 2.0)
        set_width(value_cell, 4.8)
        set_shading(label_cell, PALE_BLUE)
        for cell in (label_cell, value_cell):
            set_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_font(label_cell.paragraphs[0].add_run(str(item.get("label", ""))), 10.5, bold=True)
        value_paragraph = value_cell.paragraphs[0]
        value = str(item.get("value", ""))
        url = item.get("url")
        if url:
            add_hyperlink(value_paragraph, value or str(url), str(url))
        else:
            set_font(value_paragraph.add_run(value), 10.5)


def add_table(doc: Document, block: dict) -> None:
    headers = [str(value) for value in block.get("headers", [])]
    rows = block.get("rows", [])
    if not headers or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise ValueError("Table blocks require headers and equally sized rows.")
    widths = block.get("widths") or [6.8 / len(headers)] * len(headers)
    if len(widths) != len(headers):
        raise ValueError("Table widths must match the header count.")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    set_repeat_header(table.rows[0])
    for column, header in enumerate(headers):
        cell = table.rows[0].cells[column]
        set_width(cell, float(widths[column]))
        set_shading(cell, BLUE)
        set_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(header), 9.5, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, raw in enumerate(values):
            cell = cells[column]
            set_width(cell, float(widths[column]))
            set_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_GRAY)
            set_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_font(paragraph.add_run(str(raw)), 9.5)


def add_block(doc: Document, block: dict) -> None:
    block_type = block["type"]
    if block_type == "page_break":
        doc.add_page_break()
    elif block_type == "paragraph":
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        run = paragraph.add_run(str(block.get("text", "")))
        set_font(run, 11, bold=bool(block.get("bold")), italic=bool(block.get("italic")))
    elif block_type == "bullets":
        for item in block.get("items", []):
            paragraph = doc.add_paragraph(style="List Bullet")
            set_font(paragraph.add_run(str(item)), 11)
    elif block_type == "table":
        add_table(doc, block)
    elif block_type == "sources":
        for item in block.get("items", []):
            paragraph = doc.add_paragraph()
            label = str(item.get("label", item.get("url", "")))
            url = str(item.get("url", ""))
            prefix = str(item.get("prefix", ""))
            if prefix:
                set_font(paragraph.add_run(prefix), 10)
            if url:
                add_hyperlink(paragraph, label, url)
            else:
                set_font(paragraph.add_run(label), 10)


def build(spec: dict) -> Document:
    doc = Document()
    configure(doc, spec)
    if spec["stage"] == "DRAFT":
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.add_run("DRAFT - Not Ready for Submission"), 11, bold=True, color=RED)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run(spec["title"]), 22, bold=True)
    if spec.get("subtitle"):
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(subtitle.add_run(str(spec["subtitle"])), 13, bold=True, color=GRAY)
    add_metadata(doc, spec.get("metadata", []))
    if spec.get("opening"):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(14)
        set_font(paragraph.add_run(str(spec["opening"])), 11, bold=True)

    for section in spec["sections"]:
        if section.get("page_break_before"):
            doc.add_page_break()
        heading = section.get("heading")
        if heading:
            doc.add_heading(str(heading), level=int(section.get("level", 1)))
        for block in section.get("blocks", []):
            add_block(doc, block)
    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--dry-run", action="store_true", help="Validate the JSON spec without writing a DOCX.")
    args = parser.parse_args()
    if not args.spec.is_file():
        parser.error(f"Spec not found: {args.spec}")
    spec = load_spec(args.spec)
    if args.dry_run:
        print(f"DRY RUN PASS: {len(spec['sections'])} section(s) validated")
        return

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=args.output.stem + ".",
        suffix=".docx",
        dir=args.output.parent,
        delete=False,
    ) as temp:
        temp_path = Path(temp.name)
    try:
        build(spec).save(temp_path)
        os.replace(temp_path, args.output)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    print(f"WROTE {args.output}")


if __name__ == "__main__":
    main()
